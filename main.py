"""
AI Resume Analyzer - Primary Python Backend & CLI Service
Natural Language Processing (NLP), TF-IDF Vector Space Modeling, ATS Scoring, and Conversational AI.

Usage:
  1. Run FastAPI Web Server:
     python main.py
     # or: uvicorn main:app --reload --port 8000

  2. Run CLI Resume Analysis:
     python main.py --cli --resume path/to/resume.pdf --jd path/to/jd.txt
"""

import os
import sys
import re
import math
import uuid
import sqlite3
import argparse
from datetime import datetime
from typing import List, Optional, Dict, Any
from collections import Counter

# Third-party imports with graceful fallbacks
try:
    from fastapi import FastAPI, UploadFile, File, Form, HTTPException
    from fastapi.middleware.cors import CORSMiddleware
    from pydantic import BaseModel
    FASTAPI_AVAILABLE = True
except ImportError:
    FASTAPI_AVAILABLE = False
    FastAPI = object
    BaseModel = object

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    SKLEARN_AVAILABLE = True
except ImportError:
    SKLEARN_AVAILABLE = False

try:
    import pandas as pd
except ImportError:
    pd = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import docx
except ImportError:
    docx = None

try:
    import google.generativeai as genai
    GEMINI_AVAILABLE = True
except ImportError:
    genai = None
    GEMINI_AVAILABLE = False

# ============================================================================
# FASTAPI APP INITIALIZATION
# ============================================================================

if FASTAPI_AVAILABLE:
    app = FastAPI(
        title="AI Resume Analyzer NLP Engine (Python Main)",
        description="Academic NLP Resume Parser, TF-IDF Cosine Matcher, ATS Scorer & Gemini AI Co-Pilot",
        version="2.0.0"
    )

    app.add_middleware(
        CORSMiddleware,
        allow_origins=["*"],
        allow_credentials=True,
        allow_methods=["*"],
        allow_headers=["*"],
    )
else:
    app = None

DB_PATH = os.path.join(os.path.dirname(__file__), "resume_analyzer_sqlite.db")

# ============================================================================
# DATABASE SETUP
# ============================================================================

def init_db():
    conn = sqlite3.connect(DB_PATH)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS reports (
            id TEXT PRIMARY KEY,
            timestamp TEXT,
            candidate_name TEXT,
            job_title TEXT,
            match_percentage INTEGER,
            skills_found TEXT,
            missing_skills TEXT,
            ats_score INTEGER,
            quality_score INTEGER,
            resume_text TEXT,
            jd_text TEXT,
            strengths TEXT,
            weaknesses TEXT,
            recommendations TEXT
        )
    ''')
    conn.commit()
    conn.close()

init_db()

# ============================================================================
# NLP CORPUS & CONSTANTS
# ============================================================================

SKILLS_CORPUS = [
    # Languages
    "python", "javascript", "typescript", "java", "c++", "c#", "ruby", "go", "rust", "php", "sql", "html", "css", "r", "scala",
    # Frameworks & Libraries
    "react", "angular", "vue", "nextjs", "express", "fastapi", "django", "flask", "spring boot", "tensorflow", "pytorch",
    "scikit-learn", "pandas", "numpy", "spacy", "nltk", "tailwindcss", "bootstrap", "node.js", "graphql",
    # Databases & Cloud
    "mysql", "postgresql", "mongodb", "redis", "sqlite", "aws", "azure", "gcp", "docker", "kubernetes", "jenkins", "git",
    "ci/cd", "terraform", "firebase", "supabase",
    # Concepts
    "machine learning", "deep learning", "natural language processing", "nlp", "data science", "system design",
    "computer vision", "rest api", "microservices", "object-oriented programming", "agile", "scrum", "unit testing"
]

STOPWORDS = {
    "a", "about", "above", "after", "again", "against", "all", "am", "an", "and", "any", "are", "as", "at", "be", "because",
    "been", "before", "being", "below", "between", "both", "but", "by", "can", "could", "did", "do", "does", "doing", "down",
    "during", "each", "few", "for", "from", "further", "had", "has", "have", "having", "he", "her", "here", "hers", "him",
    "himself", "his", "how", "i", "if", "in", "into", "is", "it", "its", "itself", "me", "more", "most", "my", "myself", "no",
    "nor", "not", "of", "off", "on", "once", "only", "or", "other", "our", "ours", "ourselves", "out", "over", "own", "same",
    "she", "should", "so", "some", "such", "than", "that", "the", "their", "theirs", "them", "themselves", "then", "there",
    "these", "they", "this", "those", "through", "to", "too", "under", "until", "up", "very", "was", "we", "were", "what",
    "when", "where", "which", "while", "who", "whom", "why", "with", "you", "your", "yours", "yourself", "yourselves"
}

# ============================================================================
# NLP PREPROCESSING & ANALYSIS FUNCTIONS
# ============================================================================

def clean_text(text: str) -> str:
    """Preprocess text: lowercase, remove emails, URLs, and non-alphanumeric symbols."""
    if not text:
        return ""
    text = text.lower()
    text = re.sub(r'\S+@\S+', ' ', text)
    text = re.sub(r'http\S+', ' ', text)
    text = re.sub(r'[^a-zA-Z0-9\s\+\#\-\.]', ' ', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def tokenize(text: str) -> List[str]:
    """Tokenize clean text into non-stopword tokens."""
    cleaned = clean_text(text)
    words = cleaned.split()
    return [w for w in words if w not in STOPWORDS and len(w) > 1]

def extract_skills_from_text(text: str) -> List[str]:
    """Extract skills found in text based on skill corpus."""
    text_lower = text.lower()
    return [skill for skill in SKILLS_CORPUS if re.search(rf'\b{re.escape(skill)}\b', text_lower)]

def calculate_tf_idf_match(resume_text: str, jd_text: str) -> float:
    """Compute TF-IDF vector matrix and cosine similarity match score."""
    tokens_res = tokenize(resume_text)
    tokens_jd = tokenize(jd_text)
    
    if not tokens_res or not tokens_jd:
        return 0.0

    if SKLEARN_AVAILABLE:
        doc_res = " ".join(tokens_res)
        doc_jd = " ".join(tokens_jd)
        vectorizer = TfidfVectorizer()
        tfidf_matrix = vectorizer.fit_transform([doc_res, doc_jd])
        sim = cosine_similarity(tfidf_matrix[0:1], tfidf_matrix[1:2])[0][0]
        return float(sim)
    else:
        # Pure Python Cosine Similarity fallback
        vec_res = Counter(tokens_res)
        vec_jd = Counter(tokens_jd)
        
        intersection = set(vec_res.keys()) & set(vec_jd.keys())
        numerator = sum([vec_res[x] * vec_jd[x] for x in intersection])

        sum1 = sum([vec_res[x] ** 2 for x in vec_res.keys()])
        sum2 = sum([vec_jd[x] ** 2 for x in vec_jd.keys()])
        denominator = math.sqrt(sum1) * math.sqrt(sum2)

        if not denominator:
            return 0.0
        return float(numerator) / denominator

def extract_candidate_name(resume_text: str) -> str:
    """Extract potential candidate name from the top lines of text."""
    lines = [line.strip() for line in resume_text.split("\n") if line.strip()]
    for line in lines[:5]:
        if 3 <= len(line) <= 35 and re.match(r'^[A-Z][a-zA-Z\s\.\-]+$', line):
            if not any(kw in line.lower() for kw in ["resume", "curriculum", "vitae", "summary", "experience", "education"]):
                return line
    return "Candidate Profile"

import pickle
import json
import numpy as np

MODELS_DIR = os.path.join(os.path.dirname(__file__), "models")

def get_ml_models():
    cat_model_path = os.path.join(MODELS_DIR, "resume_category_model.pkl")
    vectorizer_path = os.path.join(MODELS_DIR, "tfidf_vectorizer.pkl")
    encoder_path = os.path.join(MODELS_DIR, "label_encoder.pkl")
    
    if not (os.path.exists(cat_model_path) and os.path.exists(vectorizer_path) and os.path.exists(encoder_path)):
        try:
            import train_model
            ds_path = os.path.join(os.path.dirname(__file__), "dataset", "resume_dataset_100.csv")
            train_model.train_pipeline(ds_path, output_dir=MODELS_DIR)
        except Exception as err:
            print("Auto training failed:", err)
            return None, None, None, None

    try:
        with open(cat_model_path, "rb") as f:
            cat_model = pickle.load(f)
        with open(vectorizer_path, "rb") as f:
            vectorizer = pickle.load(f)
        with open(encoder_path, "rb") as f:
            encoder = pickle.load(f)
        metrics = {}
        metrics_path = os.path.join(MODELS_DIR, "metrics.json")
        if os.path.exists(metrics_path):
            with open(metrics_path, "r") as f:
                metrics = json.load(f)
        return cat_model, vectorizer, encoder, metrics
    except Exception as e:
        print(f"Error loading ML models: {e}")
        return None, None, None, None

def ml_predict_resume_role(resume_text: str) -> Dict[str, Any]:
    """Perform ML Supervised Classification on candidate resume text."""
    cat_model, vectorizer, encoder, metrics = get_ml_models()
    skills = extract_skills_from_text(resume_text)
    h_res = predict_role_and_experience(skills, resume_text)
    
    if not cat_model or not vectorizer or not encoder:
        return {
            "predictedCategory": h_res["predicted_role"],
            "confidence": 75.0,
            "isMlTrained": False,
            "probabilities": {h_res["predicted_role"]: 75.0},
            "topKeywords": skills[:5],
            "experienceLevel": h_res["experience_tier"],
            "modelAccuracy": 0.0
        }
        
    cleaned = clean_text(resume_text)
    vec = vectorizer.transform([cleaned])
    probs = cat_model.predict_proba(vec)[0]
    classes = encoder.classes_
    
    prob_dict = {str(cls): float(round(p * 100, 2)) for cls, p in zip(classes, probs)}
    top_idx = probs.argmax()
    top_category = str(classes[top_idx])
    confidence = float(round(probs[top_idx] * 100, 2))
    
    feature_names = np.array(vectorizer.get_feature_names_out()) if hasattr(vectorizer, "get_feature_names_out") else np.array([])
    row_vec = vec.toarray()[0]
    top_term_indices = row_vec.argsort()[::-1][:5]
    top_keywords = [str(feature_names[i]) for i in top_term_indices if row_vec[i] > 0] if len(feature_names) > 0 else skills[:5]

    exp_model_path = os.path.join(MODELS_DIR, "experience_model.pkl")
    exp_enc_path = os.path.join(MODELS_DIR, "experience_encoder.pkl")
    exp_level = h_res["experience_tier"]
    if os.path.exists(exp_model_path) and os.path.exists(exp_enc_path):
        try:
            with open(exp_model_path, "rb") as f: exp_m = pickle.load(f)
            with open(exp_enc_path, "rb") as f: exp_e = pickle.load(f)
            exp_pred = exp_m.predict(vec)[0]
            exp_level = str(exp_e.inverse_transform([exp_pred])[0])
        except Exception:
            pass

    return {
        "predictedCategory": top_category,
        "confidence": confidence,
        "isMlTrained": True,
        "probabilities": prob_dict,
        "topKeywords": top_keywords,
        "experienceLevel": exp_level,
        "modelAccuracy": float(round(metrics.get("accuracy", 0.90) * 100, 2)),
        "datasetSize": metrics.get("dataset_records_count", 100)
    }

def predict_role_and_experience(skills: List[str], text: str) -> Dict[str, str]:
    """Predict primary job sector role and experience level."""
    skills_set = set(skills)
    
    web_score = len(skills_set.intersection({"javascript", "typescript", "react", "html", "css", "vue", "angular", "node.js"}))
    py_data_score = len(skills_set.intersection({"python", "pandas", "numpy", "scikit-learn", "tensorflow", "pytorch", "nlp", "machine learning", "data science"}))
    backend_score = len(skills_set.intersection({"java", "c++", "c#", "go", "fastapi", "django", "express", "sql", "postgresql", "mysql", "mongodb"}))
    devops_score = len(skills_set.intersection({"aws", "docker", "kubernetes", "jenkins", "gcp", "azure", "terraform", "git"}))

    scores = {
        "Full-Stack Web Developer": web_score,
        "Data Scientist / ML Engineer": py_data_score,
        "Backend Systems Engineer": backend_score,
        "DevOps & Cloud Engineer": devops_score
    }
    
    best_role = max(scores, key=scores.get)
    if scores[best_role] == 0:
        best_role = "Software Engineer"

    word_count = len(text.split())
    if word_count > 450:
        exp = "Senior Level (5+ Years)"
    elif word_count > 250:
        exp = "Mid Level (2–5 Years)"
    else:
        exp = "Entry Level / Junior (0–2 Years)"

    return {"predicted_role": best_role, "experience_tier": exp}

def extract_text_from_file_path(file_path: str) -> str:
    """Extract text from PDF, DOCX, or TXT file paths."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif ext == ".pdf":
        if pdfplumber is None:
            raise RuntimeError("pdfplumber is required for PDF parsing. Run: pip install pdfplumber")
        with pdfplumber.open(file_path) as pdf:
            return " ".join([page.extract_text() or "" for page in pdf.pages])
    elif ext == ".docx":
        if docx is None:
            raise RuntimeError("python-docx is required for DOCX parsing. Run: pip install python-docx")
        doc = docx.Document(file_path)
        return " ".join([p.text for p in doc.paragraphs])
    else:
        raise ValueError(f"Unsupported file format: {ext}")

# ============================================================================
# API SCHEMAS & ROUTES (IF FASTAPI PRESENT)
# ============================================================================

if FASTAPI_AVAILABLE:
    class AnalysisRequest(BaseModel):
        resumeText: str
        jdText: str

    class ChatRequest(BaseModel):
        message: str
        context: Optional[Dict[str, Any]] = None

    @app.get("/api/health")
    def health_check():
        return {"status": "ok", "service": "Python FastAPI NLP Engine", "timestamp": datetime.now().isoformat()}

    @app.post("/api/analyze")
    async def analyze(
        resumeFile: Optional[UploadFile] = File(None),
        resumeText: Optional[str] = Form(None),
        jdText: str = Form(...)
    ):
        text_resume = ""
        if resumeFile:
            content_bytes = await resumeFile.read()
            filename = resumeFile.filename.lower()
            if filename.endswith(".pdf") and pdfplumber:
                import io
                with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
                    text_resume = " ".join([page.extract_text() or "" for page in pdf.pages])
            elif filename.endswith(".docx") and docx:
                import io
                doc = docx.Document(io.BytesIO(content_bytes))
                text_resume = " ".join([p.text for p in doc.paragraphs])
            else:
                text_resume = content_bytes.decode("utf-8", errors="ignore")
        elif resumeText:
            text_resume = resumeText
        else:
            raise HTTPException(status_code=400, detail="Either resume file or resume text is required.")

        if not text_resume.strip():
            raise HTTPException(status_code=400, detail="Resume text is empty.")

        sim_score = calculate_tf_idf_match(text_resume, jdText)
        match_percentage = int(round(sim_score * 100))

        skills_found = extract_skills_from_text(text_resume)
        jd_skills = extract_skills_from_text(jdText)
        missing_skills = [s for s in jd_skills if s not in skills_found]

        candidate_name = extract_candidate_name(text_resume)
        role_info = predict_role_and_experience(skills_found, text_resume)

        has_email = bool(re.search(r'\S+@\S+', text_resume))
        has_phone = bool(re.search(r'\d{10,}', text_resume))
        word_count = len(text_resume.split())

        quality_score = 100
        if not has_email: quality_score -= 15
        if not has_phone: quality_score -= 15
        if word_count < 150: quality_score -= 15
        quality_score = max(30, quality_score)

        ats_score = int(round((quality_score + match_percentage) / 2))

        strengths = [
            "Structured layout with extracted technical skills section.",
            f"Predicted Candidate Role: {role_info['predicted_role']}"
        ]
        if has_email and has_phone:
            strengths.append("Complete contact details (Email & Phone detected).")

        weaknesses = []
        if missing_skills:
            weaknesses.append(f"Missing key JD skills: {', '.join(missing_skills[:3])}")
        if word_count < 150:
            weaknesses.append("Resume content length is concise/short.")

        recommendations = [
            "Incorporate highly weighted keywords from the Job Description into bullet statements.",
            "Add measurable achievements with quantifiable metrics (e.g. Improved performance by 35%).",
            "Include a concise career summary paragraph at the top of your resume."
        ]

        ml_info = ml_predict_resume_role(text_resume)
        if ml_info.get("isMlTrained"):
            role_info["predicted_role"] = ml_info["predictedCategory"]
            role_info["experience_tier"] = ml_info["experienceLevel"]
            strengths.append(f"ML Model Classification: {ml_info['predictedCategory']} ({ml_info['confidence']}% confidence)")

        report_id = f"py_rep_{uuid.uuid4().hex[:8]}"
        timestamp = datetime.now().isoformat()
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO reports VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            report_id, timestamp, candidate_name, role_info["predicted_role"], match_percentage,
            ",".join(skills_found), ",".join(missing_skills), ats_score, quality_score,
            text_resume, jdText, ",".join(strengths), ",".join(weaknesses), ",".join(recommendations)
        ))
        conn.commit()
        conn.close()

        return {
            "success": True,
            "reportId": report_id,
            "timestamp": timestamp,
            "candidateName": candidate_name,
            "jobTitle": role_info["predicted_role"],
            "matchPercentage": match_percentage,
            "atsScore": ats_score,
            "qualityScore": quality_score,
            "skillsFound": skills_found,
            "missingSkills": missing_skills,
            "strengths": strengths,
            "weaknesses": weaknesses if weaknesses else ["No critical issues found."],
            "recommendations": recommendations,
            "rolePrediction": role_info,
            "mlClassification": ml_info
        }

    class PredictRequest(BaseModel):
        resumeText: str

    @app.post("/api/predict-role")
    async def predict_role_endpoint(req: PredictRequest):
        if not req.resumeText or not req.resumeText.strip():
            raise HTTPException(status_code=400, detail="Resume text is required")
        res = ml_predict_resume_role(req.resumeText)
        return {"success": True, "data": res}

    @app.get("/api/model-info")
    def get_model_info():
        cat_m, vec_m, enc_m, metrics = get_ml_models()
        return {
            "success": True,
            "isTrained": bool(cat_m),
            "metrics": metrics
        }

    @app.post("/api/train-model")
    def trigger_train_model():
        try:
            import train_model
            ds_path = os.path.join(os.path.dirname(__file__), "dataset", "resume_dataset_100.csv")
            metrics = train_model.train_pipeline(ds_path, output_dir=MODELS_DIR)
            return {"success": True, "message": "ML Model Trained Successfully!", "metrics": metrics}
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"Training failed: {str(e)}")

    @app.post("/api/chat")
    async def chat(request: ChatRequest):
        api_key = os.environ.get("GEMINI_API_KEY")
        if GEMINI_AVAILABLE and api_key:
            try:
                genai.configure(api_key=api_key)
                model = genai.GenerativeModel("gemini-1.5-flash")
                response = model.generate_content(
                    f"You are an expert ATS Resume Coach. Answer concisely: {request.message}"
                )
                return {
                    "success": True,
                    "source": "gemini_ai",
                    "text": response.text,
                    "intent": "GENERATIVE_AI_COPILOT"
                }
            except Exception:
                pass

        return {
            "success": True,
            "source": "rule_engine",
            "text": "NLP Rule Engine Advice: Focus on adding missing technical keywords to your resume summary."
        }

    @app.get("/api/history")
    def history():
        conn = sqlite3.connect(DB_PATH)
        cursor = conn.cursor()
        cursor.execute("SELECT id, timestamp, candidate_name, job_title, match_percentage, ats_score FROM reports ORDER BY timestamp DESC")
        rows = cursor.fetchall()
        conn.close()
        
        reports = [
            {
                "id": r[0],
                "timestamp": r[1],
                "candidateName": r[2],
                "jobTitle": r[3],
                "matchPercentage": r[4],
                "atsScore": r[5]
            }
            for r in rows
        ]
        return {"success": True, "reports": reports}

# ============================================================================
# CLI RUNNER MODE
# ============================================================================

def run_cli():
    parser = argparse.ArgumentParser(description="AI Resume Analyzer - Python CLI Engine")
    parser.add_argument("--cli", action="store_true", help="Run in CLI mode")
    parser.add_argument("--resume", type=str, required=True, help="Path to resume file (.pdf, .docx, .txt)")
    parser.add_argument("--jd", type=str, required=True, help="Path to job description file (.txt)")
    args = parser.parse_args()

    print("=" * 60)
    print("🚀 AI RESUME ANALYZER - PYTHON NLP ENGINE")
    print("=" * 60)

    try:
        resume_text = extract_text_from_file_path(args.resume)
        with open(args.jd, "r", encoding="utf-8", errors="ignore") as f:
            jd_text = f.read()

        match_score = int(round(calculate_tf_idf_match(resume_text, jd_text) * 100))
        skills_found = extract_skills_from_text(resume_text)
        jd_skills = extract_skills_from_text(jd_text)
        missing_skills = [s for s in jd_skills if s not in skills_found]
        name = extract_candidate_name(resume_text)
        role = predict_role_and_experience(skills_found, resume_text)

        print(f"\n👤 Candidate Profile: {name}")
        print(f"🎯 Predicted Role:   {role['predicted_role']} ({role['experience_tier']})")
        print(f"📊 TF-IDF Cosine Match: {match_score}%")
        print(f"\n✅ Skills Found ({len(skills_found)}): {', '.join(skills_found)}")
        print(f"⚠️ Missing JD Skills ({len(missing_skills)}): {', '.join(missing_skills)}")
        print("=" * 60)

    except Exception as e:
        print(f"❌ Error during analysis: {e}", file=sys.stderr)
        sys.exit(1)

# ============================================================================
# MAIN ENTRYPOINT
# ============================================================================

if __name__ == "__main__":
    if "--cli" in sys.argv:
        run_cli()
    elif FASTAPI_AVAILABLE:
        import uvicorn
        print("Starting FastAPI AI Resume Analyzer Server on http://0.0.0.0:8000...")
        uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
    else:
        print("FastAPI is not installed. To run web server, install: pip install fastapi uvicorn")
        print("Or run in CLI mode: python main.py --cli --resume resume.txt --jd jd.txt")
